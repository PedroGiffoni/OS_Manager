from io import BytesIO
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.barcharts import VerticalBarChart

LABELS = {
    "backlog": "Backlog",
    "todo": "A fazer",
    "production": "Em produção",
    "review": "Em revisão",
    "done": "Concluído",
}


def build_export(format, rows, summary, financial, filters):
    buffer = BytesIO()
    columns = ["ID", "Título", "Status", "Previstas (min)", "Utilizadas (min)"] + (
        ["Custo aprovado (centavos)"] if financial else []
    )
    values = [
        [
            row["id"],
            row["title"],
            LABELS[row["status"]],
            row["estimated_minutes"],
            row["used_minutes"],
        ]
        + ([row["cost_cents"]] if financial else [])
        for row in rows
    ]
    formulas = [
        "Conclusão = concluídas / planejadas × 100; denominador zero resulta em 0.",
        "Pontualidade = concluídas no prazo / concluídas com prazo × 100.",
        "Desvio = minutos utilizados - minutos previstos.",
        "Período filtra a data UTC de criação da atividade; sprint encerrada usa snapshot de encerramento.",
    ]
    if financial:
        formulas.append("Custo = minutos aprovados × centavos/hora congelados / 60; ROUND_HALF_UP.")
    if format == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = "Atividades"
        sheet.append(columns + ["Fornecedor ID", "Sprint ID", "Prazo", "Aprovação"])
        for record, item in zip(values, rows):
            sheet.append(
                record + [item["supplier_id"], item["sprint_id"], item["due"], item["approval"]]
            )
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="237E85")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 28
        metrics = book.create_sheet("Indicadores")
        for key in [
            "planned",
            "done",
            "pending",
            "overdue",
            "completion_percent",
            "on_time_percent",
            "estimated_minutes",
            "used_minutes",
            "deviation_minutes",
        ]:
            metrics.append([key, str(summary[key])])
        if financial:
            metrics.append(["cost_cents", sum(row["cost_cents"] for row in rows)])
        metrics.append(["Filtros", str(filters)])
        for formula in formulas:
            metrics.append(["Definição", formula])
        metrics.column_dimensions["A"].width = 25
        metrics.column_dimensions["B"].width = 110
        book.properties.creator = "OS Manager"
        book.save(buffer)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        styles = getSampleStyleSheet()
        body = styles["BodyText"]
        body.fontSize = 8
        body.leading = 11
        story = [
            Paragraph("OS Manager · Relatório de atividades", styles["Title"]),
            Paragraph(escape(str(filters)), body),
            Spacer(1, 14),
        ]
        for formula in formulas:
            story.append(Paragraph(escape(formula), body))
        story.append(Spacer(1, 16))
        story.append(
            Paragraph(
                f"Planejadas: {summary['planned']} · Concluídas: {summary['done']} · Pendentes: {summary['pending']} · Atrasadas: {summary['overdue']} · Conclusão: {summary['completion_percent']}%",
                body,
            )
        )
        drawing = Drawing(650, 140)
        chart = VerticalBarChart()
        chart.x, chart.y, chart.width, chart.height = 40, 25, 580, 100
        chart.data = [[item["count"] for item in summary["by_status"]]]
        chart.categoryAxis.categoryNames = [LABELS[item["status"]] for item in summary["by_status"]]
        chart.valueAxis.valueMin = 0
        chart.valueAxis.valueMax = max(1, max(item["count"] for item in summary["by_status"]))
        chart.bars[0].fillColor = colors.HexColor("#237e85")
        drawing.add(chart)
        story.extend([drawing, Spacer(1, 16)])
        data = [[Paragraph(escape(str(value)), body) for value in columns]] + [
            [Paragraph(escape(str(value)), body) for value in record] for record in values
        ]
        table = Table(
            data,
            colWidths=([40, 270, 90, 95, 95, 160] if financial else [40, 350, 100, 115, 145]),
            repeatRows=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dceeed")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
                    ("TOPPADDING", (0, 0), (-1, -1), 9),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#dce5e8")),
                ]
            )
        )
        story.append(table)
        SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            leftMargin=24,
            rightMargin=24,
            topMargin=25,
            bottomMargin=25,
            title="OS Manager Relatório",
            author="OS Manager",
        ).build(story)
        media = "application/pdf"
    buffer.seek(0)
    return buffer, media
