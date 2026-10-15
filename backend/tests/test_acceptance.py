from pypdf import PdfReader
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
from openpyxl import load_workbook

P = "/api/v1"


def create(clients, supplier=1):
    response = clients["admin"].post(
        P + "/tasks", json={"title": "Teste de aceite", "supplier_id": supplier}
    )
    assert response.status_code == 201, response.text
    return response.json()


def get(clients, task, user="admin"):
    return clients[user].get(f"{P}/tasks/{task['id']}").json()


def move(clients, task, status, user="manager", **extra):
    return clients[user].post(
        f"{P}/tasks/{task['id']}/move",
        json={"version": task["version"], "status": status, **extra},
    )


def hours(clients, task, estimated, used):
    return clients["supplier"].put(
        f"{P}/tasks/{task['id']}/hours",
        json={
            "version": task["version"],
            "estimated_minutes": estimated,
            "used_minutes": used,
        },
    )


def review(clients):
    task = create(clients)
    task = move(clients, task, "todo").json()
    task = hours(clients, task, 600, 0).json()
    task = move(clients, task, "production", "supplier").json()
    task = hours(clients, task, 600, 480).json()
    return move(clients, task, "review", "supplier").json()


def decision(clients, task, approved=True, user="manager", reason=""):
    return clients[user].post(
        f"{P}/tasks/{task['id']}/approval",
        json={
            "version": task["version"],
            "hours_version": task["hours_version"],
            "approved": approved,
            "reason": reason,
        },
    )


def test_hours_transitions_and_snapshot(clients):
    task = create(clients)
    task = move(clients, task, "todo").json()
    invalid = move(clients, task, "production", "supplier")
    assert invalid.status_code == 422
    assert "horas previstas" in invalid.text
    assert get(clients, task)["status"] == "todo"
    assert get(clients, task)["version"] == task["version"]
    task = hours(clients, task, 600, 0).json()
    task = move(clients, task, "production", "supplier").json()
    assert move(clients, task, "review", "supplier").status_code == 422
    task = hours(clients, task, 600, 480).json()
    task = move(clients, task, "review", "supplier").json()
    assert clients["admin"].post(P + "/suppliers/1/rates", json={"cents": 15000}).status_code == 201
    assert get(clients, task)["cost_cents"] == 0
    assert decision(clients, task, user="no_grant").status_code == 403
    rejected = decision(clients, task, False, reason="Confira o total").json()
    assert rejected["status"] == "review" and rejected["approval"] == "rejected"
    task = hours(clients, rejected, 600, 480).json()
    assert task["approval"] == "resubmitted"
    approved_response = decision(clients, task)
    assert approved_response.status_code == 200, approved_response.text
    approved = approved_response.json()
    assert approved["cost_cents"] == 96000
    assert decision(clients, task).status_code == 200
    assert move(clients, approved, "done", "supplier").status_code == 403
    done = move(clients, approved, "done").json()
    assert done["status"] == "done"
    reopened = move(clients, done, "review", reason="Correção autorizada").json()
    assert reopened["cost_cents"] == 0 and reopened["approval"] == "pending"
    assert hours(clients, reopened, 600, 420).status_code == 200


def test_horizontal_financial_exports_and_csrf(clients):
    task = create(clients)
    other = create(clients, 2)
    payload = clients["supplier"].get(P + "/tasks").json()
    assert [row["id"] for row in payload] == [task["id"]]
    assert "rate_cents" not in str(payload) and "cost_cents" not in str(payload)
    assert clients["supplier"].get(f"{P}/tasks/{other['id']}").status_code == 404
    assert clients["supplier"].get(P + "/suppliers/1/rates").status_code == 403
    assert clients["supplier"].get(P + "/exports/xlsx?financial=true").status_code == 403
    assert "cost" not in clients["supplier"].get(P + "/dashboard").text
    response = clients["supplier"].get(P + "/exports/xlsx")
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.max_row == 2
    assert "Custo" not in str(list(sheet.values))
    client = clients["admin"]
    token = client.headers.pop("X-CSRF-Token")
    assert client.post(P + "/tasks", json={"title": "CSRF"}).status_code == 403
    client.headers["X-CSRF-Token"] = token
    assert (
        clients["manager"]
        .put(
            f"{P}/tasks/{task['id']}/hours",
            json={
                "version": task["version"],
                "estimated_minutes": 60,
                "used_minutes": 0,
            },
        )
        .status_code
        == 403
    )


def test_sprint_close_preserves_history_and_uniqueness(clients):
    admin = clients["admin"]
    task = create(clients)
    ids = [
        admin.post(
            P + "/sprints",
            json={
                "name": name,
                "objective": "Aceite",
                "starts": "2026-01-01",
                "ends": "2026-01-15",
            },
        ).json()["id"]
        for name in ["A", "B"]
    ]
    included = admin.post(
        f"{P}/sprints/{ids[0]}/tasks",
        json={"version": 1, "task_id": task["id"], "task_version": task["version"]},
    )
    assert included.status_code == 200, included.text
    task = included.json()
    assert (
        admin.post(
            f"{P}/sprints/{ids[1]}/tasks",
            json={"version": 1, "task_id": task["id"], "task_version": task["version"]},
        ).status_code
        == 409
    )
    assert (
        admin.post(
            f"{P}/sprints/{ids[0]}/action", json={"version": 2, "action": "start"}
        ).status_code
        == 200
    )
    assert (
        admin.post(
            f"{P}/sprints/{ids[0]}/action", json={"version": 3, "action": "close"}
        ).status_code
        == 200
    )
    returned = get(clients, task)
    assert returned["sprint_id"] is None and returned["status"] == "backlog"
    assert "returned_to_backlog" in admin.get(f"{P}/audit?entity=task&identity={task['id']}").text


def test_card_schema_preserves_values_and_native_rules(clients):
    task = create(clients)
    field = {
        "id": "evidence",
        "label": "Evidência",
        "type": "short_text",
        "required_for": ["production"],
    }
    response = clients["admin"].put(
        f"{P}/tasks/{task['id']}/schema",
        json={
            "version": task["version"],
            "fields": [field],
            "values": {"evidence": "Documento"},
        },
    )
    assert response.status_code == 200, response.text
    task = response.json()
    assert (
        clients["admin"]
        .put(
            f"{P}/tasks/{task['id']}/schema",
            json={"version": task["version"], "fields": [], "values": {}},
        )
        .status_code
        == 422
    )
    task = move(clients, task, "todo").json()
    assert move(clients, task, "production").status_code == 422


def test_stale_drag_and_concurrent_approval(clients):
    task = create(clients)
    assert move(clients, task, "todo").status_code == 200
    assert move(clients, task, "todo").status_code == 409
    task = review(clients)
    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: decision(clients, task).status_code, range(2)))
    assert all(result in {200, 409} for result in results), results
    assert 200 in results
    assert get(clients, task)["cost_cents"] == 96000


def test_exports_filters_and_financial_scope(clients):
    task = review(clients)
    task = decision(clients, task).json()
    create(clients, 2)
    admin = clients["admin"]
    xlsx = admin.get(P + "/exports/xlsx?financial=true&supplier_id=1&status=review")
    sheet = load_workbook(BytesIO(xlsx.content)).active
    assert sheet.max_row == 2 and sheet.cell(2, 6).value == 96000
    pdf = admin.get(P + "/exports/pdf?financial=true&supplier_id=1")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert clients["no_grant"].get(P + "/exports/xlsx").status_code == 403


def test_pdf_content_respects_financial_permission(clients):
    task = review(clients)
    decision(clients, task)
    financial = clients["admin"].get(P + "/exports/pdf?financial=true")
    text = " ".join(page.extract_text() for page in PdfReader(BytesIO(financial.content)).pages)
    assert "96000" in text and "Custo aprovado" in text
    operational = clients["supplier"].get(P + "/exports/pdf")
    text = " ".join(page.extract_text() for page in PdfReader(BytesIO(operational.content)).pages)
    assert "96000" not in text and "Custo aprovado" not in text and "centavos/hora" not in text
