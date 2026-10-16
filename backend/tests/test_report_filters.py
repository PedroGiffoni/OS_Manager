import json
from io import BytesIO
from openpyxl import load_workbook
from pypdf import PdfReader
from .test_acceptance import P, create, get, hours, review, decision


def condition(field, op="eq", value="", value_to=""):
    return {"field": field, "op": op, "value": str(value), "value_to": str(value_to)}


def filtered(client, path, conditions, **params):
    return client.get(P + path, params={"field_filters": json.dumps(conditions), **params})


def schema(client, task, fields, values=None):
    response = client.put(
        f"{P}/tasks/{task['id']}/schema",
        json={"version": task["version"], "fields": fields, "values": values or {}},
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_all_card_field_types_and_combined_export_filters(clients):
    admin = clients["admin"]
    fields = [
        {"id": "short", "label": "Unidade", "type": "short_text"},
        {"id": "long", "label": "Observações", "type": "long_text"},
        {"id": "number", "label": "Quantidade", "type": "number"},
        {"id": "date", "label": "Entrega", "type": "date"},
        {"id": "select", "label": "Região", "type": "select", "options": ["Sul", "Norte"]},
        {"id": "checkbox", "label": "Prioritário", "type": "checkbox"},
        {"id": "user", "label": "Revisor", "type": "user"},
        {"id": "supplier", "label": "Profissional", "type": "supplier"},
        {"id": "attachment", "label": "Comprovante", "type": "attachment"},
        {"id": "empty", "label": "Sem resposta", "type": "short_text"},
    ]
    first, second = create(clients), create(clients)
    uploaded = admin.post(
        f"{P}/tasks/{first['id']}/attachments",
        files={"file": ("evidencia.txt", b"teste", "text/plain")},
    )
    assert uploaded.status_code == 201, uploaded.text
    first = get(clients, first)
    values = {
        "short": "Brasília",
        "long": "Atendimento presencial",
        "number": "1.25",
        "date": "2026-10-10",
        "select": "Sul",
        "checkbox": False,
        "user": 1,
        "supplier": 1,
        "attachment": [uploaded.json()["id"]],
    }
    first = schema(admin, first, fields, values)
    second = schema(
        admin,
        second,
        fields,
        {
            **values,
            "short": "Recife",
            "number": "9",
            "date": "2026-11-10",
            "select": "Norte",
            "checkbox": True,
            "attachment": [],
        },
    )
    first = hours(clients, first, 90, 0).json()
    definitions = admin.get(P + "/reports/fields").json()
    keys = {item["key"] for item in definitions}
    assert "native:description" not in keys
    assert all("field:" + item["id"] in keys for item in fields)
    cases = [
        condition("native:id", value=first["id"]),
        condition("native:estimated_minutes", "between", "1,5", "2"),
        condition("native:priority", value="0"),
        condition("native:creator_id", value="1"),
        condition("field:short", "contains", "BRASÍLIA"),
        condition("field:long", "contains", "presencial"),
        condition("field:number", "between", "1.20", "1.30"),
        condition("field:date", "between", "2026-10-01", "2026-10-20"),
        condition("field:select", value="Sul"),
        condition("field:checkbox", value="false"),
        condition("field:user", value="1"),
        condition("field:supplier", value="1"),
        condition("field:attachment", "contains", "evidencia"),
        condition("native:attachments", "contains", "EVIDENCIA.TXT"),
    ]
    for item in cases:
        response = filtered(admin, "/reports/preview", [item])
        assert response.status_code == 200, response.text
        ids = [row["id"] for row in response.json()["items"]]
        assert first["id"] in ids, (item, response.text)
        if item["field"] not in {
            "native:priority",
            "native:creator_id",
            "field:long",
            "field:user",
            "field:supplier",
        }:
            assert second["id"] not in ids, (item, response.text)
    assert (
        filtered(admin, "/reports/preview", [condition("field:empty", "empty")]).json()["total"]
        == 2
    )
    assert (
        filtered(admin, "/reports/preview", [condition("field:checkbox", value="false")]).json()[
            "total"
        ]
        == 1
    )
    combined = [condition("field:select", value="Sul"), condition("field:number", "lte", "2")]
    assert filtered(admin, "/dashboard", combined).json()["planned"] == 1
    xlsx = filtered(admin, "/exports/xlsx", combined)
    assert xlsx.status_code == 200
    book = load_workbook(BytesIO(xlsx.content))
    assert book.active.max_row == 2 and book.active.cell(2, 1).value == first["id"]
    assert "Região" in str(list(book["Indicadores"].values))
    pdf = filtered(admin, "/exports/pdf", combined)
    assert pdf.status_code == 200
    assert "Região" in "".join(
        page.extract_text() for page in PdfReader(BytesIO(pdf.content)).pages
    )
    impossible = [condition("field:select", value="Sul"), condition("field:number", "gte", "8")]
    assert filtered(admin, "/reports/preview", impossible).json()["total"] == 0
    assert filtered(admin, "/dashboard", impossible).json()["completion_percent"] == "0"
    assert filtered(admin, "/reports/preview", combined, offset=1, limit=1).json()["items"] == []


def test_report_filter_scope_and_hidden_fields(clients):
    admin, supplier = clients["admin"], clients["supplier"]
    task, other = create(clients), create(clients, 2)
    hidden = {
        "id": "private",
        "label": "Segredo interno",
        "type": "short_text",
        "visible_to": ["admin", "manager"],
    }
    schema(admin, task, [hidden], {"private": "Reservado"})
    schema(
        admin,
        other,
        [{"id": "other", "label": "Somente outro fornecedor", "type": "short_text"}],
        {"other": "Outro"},
    )
    catalogue = supplier.get(P + "/reports/fields")
    assert catalogue.status_code == 200
    assert all(
        value not in catalogue.text
        for value in [
            "Segredo interno",
            "Somente outro fornecedor",
            "rate_cents",
            "cost_cents",
            "Reservado",
        ]
    )
    for field in [
        "field:private",
        "field:other",
        "native:rate_cents",
        "native:cost_cents",
        "native:description",
    ]:
        for path in ["/reports/preview", "/dashboard", "/exports/xlsx", "/exports/pdf"]:
            result = filtered(supplier, path, [condition(field, "not_empty")])
            assert result.status_code == 403, result.text
            assert "Reservado" not in result.text
    assert (
        filtered(supplier, "/reports/preview", [condition("native:id", value=other["id"])]).json()[
            "total"
        ]
        == 0
    )
    assert filtered(supplier, "/reports/preview", []).json()["total"] == 1
    no_grant = clients["no_grant"].get(P + "/reports/fields")
    assert "cost_cents" not in no_grant.text


def test_invalid_filter_types_ranges_and_limits(clients):
    admin = clients["admin"]
    create(clients)
    invalid = [
        [condition("native:estimated_minutes", "between", "8", "2")],
        [condition("native:estimated_minutes", "eq", "NaN")],
        [condition("native:estimated_minutes", "eq", "Infinity")],
        [condition("native:due", "eq", "2026-02-30")],
        [condition("native:id", "contains", "1")],
        [condition("native:status", "eq", "invalid")],
        [condition("native:supplier_id", "eq", "abc")],
        [condition("native:title", "eq", "")],
        [condition("native:title", "contains", "x")] * 31,
    ]
    for value in invalid:
        response = filtered(admin, "/reports/preview", value)
        assert response.status_code == 422, response.text
    for raw in ["{", "{}", "null", '[{"field":"native:id","op":"eq","value":true}]']:
        assert admin.get(P + "/reports/preview", params={"field_filters": raw}).status_code == 422


def test_financial_filters_need_financial_access_and_export_permission(clients):
    task = decision(clients, review(clients)).json()
    condition_cost = [condition("native:cost_cents", "between", "960", "960")]
    assert filtered(clients["manager"], "/reports/preview", condition_cost).json()["total"] == 1
    assert filtered(clients["admin"], "/exports/xlsx", condition_cost).status_code == 200
    clients["admin"].put(P + "/users/2/grants", json=["view_financial", "export_data"])
    assert filtered(clients["manager"], "/reports/preview", condition_cost).status_code == 200
    assert filtered(clients["manager"], "/exports/xlsx", condition_cost).status_code == 403
    assert (
        filtered(
            clients["manager"], "/reports/preview", [condition("native:id", value=task["id"])]
        ).json()["total"]
        == 1
    )


def test_closed_sprint_filters_use_historical_card_schema_and_values(clients):
    admin = clients["admin"]
    task = create(clients)
    fields = [{"id": "unit", "label": "Unidade histórica", "type": "short_text"}]
    task = schema(admin, task, fields, {"unit": "Antes"})
    sprint = admin.post(
        P + "/sprints",
        json={
            "name": "Histórica",
            "objective": "Teste",
            "starts": "2026-01-01",
            "ends": "2026-01-15",
        },
    ).json()
    task = admin.post(
        f"{P}/sprints/{sprint['id']}/tasks",
        json={"version": 1, "task_id": task["id"], "task_version": task["version"]},
    ).json()
    for action, version in [("start", 2), ("close", 3)]:
        response = admin.post(
            f"{P}/sprints/{sprint['id']}/action", json={"version": version, "action": action}
        )
        assert response.status_code == 200, response.text
    task = get(clients, task)
    schema(admin, task, fields, {"unit": "Depois"})
    before = [condition("field:unit", "eq", "Antes")]
    assert filtered(admin, "/reports/preview", before, sprint_id=sprint["id"]).json()["total"] == 1
    assert filtered(admin, "/reports/preview", before).json()["total"] == 0
    assert filtered(admin, "/exports/xlsx", before, sprint_id=sprint["id"]).status_code == 200
