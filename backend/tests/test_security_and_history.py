from .test_acceptance import hours
import os
import subprocess
import sys
from io import BytesIO
from openpyxl import load_workbook
from .test_acceptance import P, create, get, move, review, decision


def test_attachment_scope_validation_and_comment_history(clients):
    task = create(clients)
    supplier = clients["supplier"]
    response = supplier.post(
        f"{P}/tasks/{task['id']}/attachments",
        files={"file": ("evidence.txt", b"Fictitious evidence", "text/plain")},
    )
    assert response.status_code == 201, response.text
    identity = response.json()["id"]
    assert supplier.get(f"{P}/attachments/{identity}").content == b"Fictitious evidence"
    assert clients["other"].get(f"{P}/attachments/{identity}").status_code == 404
    assert (
        supplier.post(
            f"{P}/tasks/{task['id']}/attachments",
            files={"file": ("unsafe.html", b"<script>alert(1)</script>", "text/html")},
        ).status_code
        == 422
    )
    assert (
        supplier.post(
            f"{P}/tasks/{task['id']}/comments", json={"body": "Comentário fictício"}
        ).status_code
        == 201
    )
    assert (
        supplier.get(f"{P}/tasks/{task['id']}/comments").json()[0]["body"] == "Comentário fictício"
    )
    assert clients["other"].get(f"{P}/tasks/{task['id']}/comments").status_code == 404


def test_closed_sprint_snapshots_and_formal_reopen(clients):
    task = review(clients)
    task = decision(clients, task).json()
    task = move(clients, task, "done").json()
    # A completed unplanned task cannot be silently assigned to a sprint.
    sprint = (
        clients["admin"]
        .post(
            P + "/sprints",
            json={
                "name": "Histórico",
                "objective": "Snapshot",
                "starts": "2026-01-01",
                "ends": "2026-02-01",
            },
        )
        .json()["id"]
    )
    assert (
        clients["admin"]
        .post(
            f"{P}/sprints/{sprint}/tasks",
            json={"version": 1, "task_id": task["id"], "task_version": task["version"]},
        )
        .status_code
        == 409
    )
    pending = create(clients)
    assert (
        clients["admin"]
        .post(
            f"{P}/sprints/{sprint}/tasks",
            json={"version": 1, "task_id": pending["id"], "task_version": pending["version"]},
        )
        .status_code
        == 200
    )
    clients["admin"].post(f"{P}/sprints/{sprint}/action", json={"version": 2, "action": "start"})
    clients["admin"].post(f"{P}/sprints/{sprint}/action", json={"version": 3, "action": "close"})
    summary = clients["admin"].get(f"{P}/dashboard?sprint_id={sprint}").json()
    assert summary["planned"] == 1 and summary["pending"] == 1
    returned = get(clients, pending)
    assert returned["status"] == "backlog"
    clients["admin"].put(
        f"{P}/tasks/{pending['id']}",
        json={
            "title": "Outro título",
            "description": "Alteração posterior",
            "supplier_id": 1,
            "version": returned["version"],
            "priority": 0,
        },
    )
    assert clients["admin"].get(f"{P}/dashboard?sprint_id={sprint}").json()["planned"] == 1
    assert (
        clients["no_grant"]
        .post(
            f"{P}/sprints/{sprint}/action",
            json={"version": 4, "action": "reopen", "reason": "Teste"},
        )
        .status_code
        == 403
    )


def test_active_supplier_requires_atomic_rate_and_live_permission_revocation(clients):
    admin = clients["admin"]
    assert (
        admin.post(
            P + "/suppliers", json={"name": "Sem tarifa", "function": "Teste", "active": True}
        ).status_code
        == 422
    )
    response = admin.post(
        P + "/suppliers",
        json={
            "name": "Com tarifa",
            "function": "Teste",
            "active": True,
            "initial_rate_cents": 12000,
        },
    )
    assert response.status_code == 201, response.text
    identity = response.json()["id"]
    assert admin.get(f"{P}/suppliers/{identity}/rates").json()[0]["cents"] == 12000
    assert clients["manager"].get(P + "/dashboard").json().get("cost_cents") == 0
    assert admin.put(P + "/users/2/grants", json=[]).status_code == 200
    assert "cost_cents" not in clients["manager"].get(P + "/dashboard").json()
    task = review(clients)
    assert decision(clients, task).status_code == 403


def test_archive_preserves_records_and_spreadsheet_formulas_are_text(clients):
    task = create(clients)
    task = (
        clients["admin"]
        .put(
            f"{P}/tasks/{task['id']}",
            json={
                "title": '=HYPERLINK("http://example.invalid")',
                "supplier_id": 1,
                "version": task["version"],
            },
        )
        .json()
    )
    exported = clients["admin"].get(P + "/exports/xlsx")
    cell = load_workbook(BytesIO(exported.content)).active.cell(2, 2)
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")
    assert (
        clients["admin"]
        .delete(f"{P}/tasks/{task['id']}?expected_version={task['version']}")
        .status_code
        == 200
    )
    assert clients["admin"].get(f"{P}/tasks/{task['id']}").status_code == 404
    assert "archived" in clients["admin"].get(f"{P}/audit?entity=task&identity={task['id']}").text


def test_reorder_atomic_versions_and_no_cross_supplier_access(clients):
    first = create(clients)
    second = create(clients)
    response = clients["supplier"].post(
        P + "/tasks/reorder",
        json=[
            {"id": second["id"], "version": second["version"]},
            {"id": first["id"], "version": first["version"]},
        ],
    )
    assert response.status_code == 200, response.text
    assert [row["id"] for row in clients["supplier"].get(P + "/tasks").json()] == [
        second["id"],
        first["id"],
    ]
    assert (
        clients["supplier"]
        .post(P + "/tasks/reorder", json=[{"id": first["id"], "version": first["version"]}])
        .status_code
        == 409
    )
    foreign = create(clients, 2)
    assert (
        clients["supplier"]
        .post(P + "/tasks/reorder", json=[{"id": foreign["id"], "version": foreign["version"]}])
        .status_code
        == 404
    )


def test_migrations_upgrade_downgrade_empty_database(tmp_path):
    env = {**os.environ, "DATABASE_URL": "sqlite:///" + str(tmp_path / "migration.db")}
    for action, revision in [("upgrade", "head"), ("downgrade", "base"), ("upgrade", "head")]:
        completed = subprocess.run(
            [sys.executable, "-m", "alembic", action, revision],
            env=env,
            capture_output=True,
            text=True,
        )
        assert completed.returncode == 0, completed.stderr


def test_schema_endpoint_cannot_bypass_value_permissions(clients):
    task = create(clients)
    field = {
        "id": "restricted",
        "label": "Campo restrito",
        "type": "short_text",
        "visible_to": ["admin", "manager"],
        "editable_by": ["manager"],
        "required_for": [],
    }
    response = clients["admin"].put(
        f"{P}/tasks/{task['id']}/schema",
        json={"version": task["version"], "fields": [field], "values": {}},
    )
    assert response.status_code == 200, response.text
    task = response.json()
    response = clients["manager"].put(
        f"{P}/tasks/{task['id']}/values",
        json={"version": task["version"], "values": {"restricted": "Valor protegido"}},
    )
    assert response.status_code == 200, response.text
    task = response.json()
    assert clients["supplier"].get(f"{P}/tasks/{task['id']}").json()["fields"] == []
    assert (
        clients["admin"]
        .put(
            f"{P}/tasks/{task['id']}/schema",
            json={
                "version": task["version"],
                "fields": [field],
                "values": {"restricted": "Tentativa"},
            },
        )
        .status_code
        == 403
    )
    assert (
        clients["supplier"]
        .put(
            f"{P}/tasks/{task['id']}/values",
            json={"version": task["version"], "values": {"restricted": "Tentativa"}},
        )
        .status_code
        == 403
    )
    assert "access_denied" in clients["admin"].get(P + "/audit?entity=security&identity=4").text


def test_monetary_half_up_uses_approved_minutes_not_estimate(clients):
    assert clients["admin"].post(P + "/suppliers/1/rates", json={"cents": 30}).status_code == 201
    task = create(clients)
    task = move(clients, task, "todo").json()
    task = hours(clients, task, 600, 0).json()
    task = move(clients, task, "production", "supplier").json()
    task = hours(clients, task, 600, 1).json()
    task = move(clients, task, "review", "supplier").json()
    assert get(clients, task)["cost_cents"] == 0
    task = decision(clients, task).json()
    assert task["cost_cents"] == 1  # 1 minute * 30 cents/hour / 60 = 0.5 cent -> 1 cent


def test_native_fields_reference_structural_hours_and_cannot_reference_rates(clients):
    task = create(clients)
    fields = [
        {
            "id": "native_estimate",
            "label": "Planejado nativo",
            "type": "native",
            "native_key": "estimated_minutes",
            "editable_by": ["supplier"],
            "required_for": ["production"],
        },
        {
            "id": "native_used",
            "label": "Realizado nativo",
            "type": "native",
            "native_key": "used_minutes",
            "editable_by": ["supplier"],
            "required_for": ["review"],
        },
    ]
    response = clients["admin"].put(
        f"{P}/tasks/{task['id']}/schema",
        json={"version": task["version"], "fields": fields, "values": {}},
    )
    assert response.status_code == 200, response.text
    task = response.json()
    assert (
        clients["supplier"]
        .put(
            f"{P}/tasks/{task['id']}/values",
            json={"version": task["version"], "values": {"native_estimate": 600}},
        )
        .status_code
        == 422
    )
    malicious = [{"id": "tariff", "label": "Tarifa", "type": "native", "native_key": "rate_cents"}]
    assert (
        clients["admin"]
        .put(
            f"{P}/tasks/{task['id']}/schema",
            json={"version": task["version"], "fields": malicious, "values": {}},
        )
        .status_code
        == 422
    )
    task = move(clients, task, "todo").json()
    task = hours(clients, task, 600, 0).json()
    task = move(clients, task, "production", "supplier").json()
    assert task["status"] == "production", task
    task = hours(clients, task, 600, 480).json()
    task = move(clients, task, "review", "supplier").json()
    assert task["status"] == "review", task
    assert task["values"]["native_used"] == 480
    task = decision(clients, task).json()
    assert task["cost_cents"] == 96000


def test_replanning_approved_hours_does_not_double_attribute_sprint_cost(clients):
    admin = clients["admin"]
    ids = [
        admin.post(
            P + "/sprints",
            json={
                "name": name,
                "objective": "Custo por aprovação",
                "starts": "2026-01-01",
                "ends": "2026-01-15",
            },
        ).json()["id"]
        for name in ["Origem", "Destino"]
    ]
    task = create(clients)
    response = admin.post(
        f"{P}/sprints/{ids[0]}/tasks",
        json={"version": 1, "task_id": task["id"], "task_version": task["version"]},
    )
    task = response.json()
    assert (
        admin.post(
            f"{P}/sprints/{ids[0]}/action", json={"version": 2, "action": "start"}
        ).status_code
        == 200
    )
    task = hours(clients, task, 600, 0).json()
    task = move(clients, task, "production", "supplier").json()
    task = hours(clients, task, 600, 480).json()
    task = move(clients, task, "review", "supplier").json()
    task = decision(clients, task).json()
    assert task["cost_cents"] == 96000
    assert (
        admin.post(
            f"{P}/sprints/{ids[0]}/action", json={"version": 3, "action": "close"}
        ).status_code
        == 200
    )
    task = get(clients, task)
    assert task["status"] == "backlog" and task["approval"] == "approved"
    assert (
        admin.post(
            f"{P}/sprints/{ids[1]}/tasks",
            json={"version": 1, "task_id": task["id"], "task_version": task["version"]},
        ).status_code
        == 200
    )
    assert admin.get(f"{P}/dashboard?sprint_id={ids[0]}").json()["cost_cents"] == 96000
    assert admin.get(f"{P}/dashboard?sprint_id={ids[1]}").json()["cost_cents"] == 0
    assert admin.get(P + "/dashboard").json()["cost_cents"] == 96000
